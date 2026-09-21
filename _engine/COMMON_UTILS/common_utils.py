import time
import requests
import json
import datetime
from pytz import timezone
import pandas as pd
import os
from requests.adapters import HTTPAdapter

# Retry budget: was 60 attempts x 5s = 300s of silent sleeping on every bad call,
# after which the function fell off the end returning None (caller then hit
# AttributeError on .json()). Now bounded, and it raises with the real cause.
max_retries = 5
retry_delay = 5
ES_TIMEOUT = (5, 180)            # (connect, read) - previously NO timeout at all
RETRY_STATUS = {429, 502, 503, 504}

# One pooled session for every call made through get_resp. Previously each call used
# a bare requests.post, i.e. a fresh TCP+TLS handshake per request.
_HTTP = requests.Session()
_HTTP.mount("https://", HTTPAdapter(pool_connections=32, pool_maxsize=32))
_HTTP.mount("http://", HTTPAdapter(pool_connections=32, pool_maxsize=32))

# --- Elasticsearch authentication -------------------------------------------
# The credential is NEVER stored in this repository. It is read from the
# environment at run time, so this file is safe to publish:
#
#     export ELASTIC_AUTH='Basic <base64 of "username:password">'
#
# See README.md -> Configuration.
ELASTIC_ENCRYPTED_PASSWORD = os.environ.get("ELASTIC_AUTH", "").strip()
if not ELASTIC_ENCRYPTED_PASSWORD:
    raise RuntimeError(
        "ELASTIC_AUTH is not set - refusing to run without Elasticsearch credentials.\n"
        "    export ELASTIC_AUTH='Basic <base64 of \"username:password\">'\n"
        "See README.md (Configuration)."
    )
def _active_cycles():
    v = os.environ.get("CYCLES", "").strip()
    return [c for c in v.split(",") if c] if v else []



CYCLE_INDEXES = ("project-task-index", "hf-referral-index", "referral-index",
                 "stock-index", "crosscycle-index")


def cycle_terms(cycles):
    """Cycle values plus their padded/unpadded variants.

    Five tenants (go, jg, ko, na, oy) carry a trace of unpadded values ('1' where
    every other doc says '01'), so a bare term match would miss them.
    """
    out = set()
    for c in cycles or []:
        c = str(c).strip()
        if not c:
            continue
        out.add(c)
        if c.isdigit():
            out.add(str(int(c)))
            out.add("%02d" % int(c))
    return sorted(out)


def get_resp(url, data, es=False):
    cyc = _active_cycles()
    if es and cyc and any(t in url for t in CYCLE_INDEXES) and isinstance(data, dict):
        try:
            q = data.get("query")
            key = "Data.additionalDetails.cycleIndex.keyword"
            if isinstance(q, dict) and isinstance(q.get("bool"), dict):
                b = q["bool"]
                slot = "must" if isinstance(b.get("must"), list) else ("filter" if isinstance(b.get("filter"), list) else None)
                if slot is not None:
                    data = json.loads(json.dumps(data))  # copy; avoid mutating caller/duplicating
                    arr = data["query"]["bool"][slot]
                    # check BOTH "terms" and "term": a script that keeps its own
                    # single-value {"term": {...}} clause must not get a second,
                    # possibly conflicting, clause appended (that ANDs to zero rows).
                    present = any(isinstance(m, dict) and
                                  (key in m.get("terms", {}) or key in m.get("term", {}))
                                  for m in arr)
                    if not present:
                        arr.append({"terms": {key: cycle_terms(cyc)}})
        except Exception:
            pass
    headers = {"Content-Type": "application/json"}
    if es:
        headers["Authorization"] = ELASTIC_ENCRYPTED_PASSWORD
    body = json.dumps(data)
    last = None
    for attempt in range(1, max_retries + 1):
        try:
            response = _HTTP.post(url, data=body, headers=headers,
                                  verify=False, timeout=ES_TIMEOUT)
            if response.status_code == 200:
                return response
            last = "HTTP %s: %s" % (response.status_code, (response.text or "")[:300])
            if response.status_code not in RETRY_STATUS:
                # 4xx (bad query, missing index, auth) will never succeed on retry
                raise RuntimeError("get_resp: non-retryable %s for %s" % (last, url))
        except requests.exceptions.RequestException as ex:
            last = "%s: %s" % (type(ex).__name__, ex)
        if attempt < max_retries:
            print("  [get_resp] attempt %d/%d failed (%s); retry in %ds"
                  % (attempt, max_retries, last, retry_delay), flush=True)
            time.sleep(retry_delay)
    raise RuntimeError("get_resp: gave up after %d attempts on %s - last error: %s"
                       % (max_retries, url, last))

def epoch_millis_to_cat(milli_seconds):
    seconds = milli_seconds / 1000.0
    utc_dt = datetime.datetime.fromtimestamp(seconds, tz=datetime.timezone.utc)
    cat_tz = timezone('Africa/Harare')
    cat_dt = utc_dt.astimezone(cat_tz)
    formatted_time = cat_dt.strftime('%d-%m-%y %H:%M:%S')
    return formatted_time


def current_date_time_in_cat():
    cat_tz = timezone('Africa/Harare')
    cat_time = datetime.datetime.now(cat_tz)
    return cat_time.strftime('%Y-%m-%d %H:%M:%S %Z')


# def load_query(query_name):
#     with open(PATH_TO_QUERIES_JSON, "r") as queries:
#         all_queries = json.load(queries)
#     return all_queries.get(query_name)


def save_excel(df, path, max_rows=1_000_000, extra_sheets=None):
    """Write df to .xlsx.

    SPLIT RULE (applies to EVERY report): if len(df) > max_rows the output is
    split into SEPARATE FILES -- <base>_part1.xlsx, <base>_part2.xlsx, ... --
    each holding at most max_rows rows. At or below max_rows a single <base>.xlsx
    is written. Rows are streamed (never sliced into new DataFrames) so peak RAM
    does not grow with the number of parts.

    extra_sheets ({name: DataFrame}) are written into the FIRST file only.

    Engine: xlsxwriter constant_memory (row-streamed, low RAM) if available,
    else openpyxl write_only. There is deliberately NO CSV fallback -- if no xlsx
    engine is installed this raises rather than silently changing the format.

    Returns the path written, or a list of paths when the output was split.
    """
    import math, os
    _d = os.path.dirname(os.path.abspath(path))
    if _d: os.makedirs(_d, exist_ok=True)
    n = len(df)
    nparts = max(1, math.ceil(n / max_rows)) if n else 1
    cols = [str(c) for c in df.columns]
    base, ext = os.path.splitext(path)
    if not ext: ext = ".xlsx"

    # Per-cell cost used to be: isinstance() + hasattr() on EVERY value
    # (rows x cols calls - ~1.8M for a 73k x 25 frame). Only numpy-backed columns
    # can yield a scalar needing .item(), so precompute which positions those are
    # and skip both checks for the object/string columns (the majority here).
    _needs_item = [dt.kind in "iufbmM" for dt in df.dtypes]

    def _row(row, _ni=_needs_item):
        out = []
        ap = out.append
        for v, ni in zip(row, _ni):
            if v is None or v != v:      # None or NaN; cheaper than isinstance
                ap(None)
            elif ni:
                ap(v.item() if hasattr(v, "item") else v)
            else:
                ap(v)
        return out

    engine = None
    try:
        import xlsxwriter  # noqa: F401
        engine = "xlsxwriter"
    except ImportError:
        try:
            from openpyxl import Workbook  # noqa: F401
            engine = "openpyxl"
        except ImportError:
            engine = None
    if engine is None:
        raise RuntimeError(
            "save_excel: no xlsx engine available (need xlsxwriter or openpyxl). "
            "Refusing to fall back to CSV. Fix with: pip install --user xlsxwriter")

    targets = ([path] if nparts == 1
               else [f"{base}_part{i}{ext}" for i in range(1, nparts + 1)])
    it = df.itertuples(index=False, name=None)
    written = []

    for pi, target in enumerate(targets, start=1):
        extras = extra_sheets if pi == 1 else None
        rows = 0
        if engine == "xlsxwriter":
            import xlsxwriter
            wb = xlsxwriter.Workbook(target, {"constant_memory": True,
                                              "strings_to_urls": False})
            ws = wb.add_worksheet("Sheet1"); ws.write_row(0, 0, cols)
            r = 1
            for _ in range(max_rows):
                try: row = next(it)
                except StopIteration: break
                ws.write_row(r, 0, _row(row)); r += 1; rows += 1
                # heartbeat: this loop is the single longest stage of a big report and
                # was previously completely silent, so the UI label looked frozen.
                if rows % 25000 == 0:
                    print(f"  [save_excel] {rows:,}/{n:,} rows written "
                          f"({100.0 * rows / n:.0f}%) -> {os.path.basename(target)}",
                          flush=True)
            for nm, edf in (extras or {}).items():
                ews = wb.add_worksheet(str(nm)[:31])
                ews.write_row(0, 0, [str(c) for c in edf.columns])
                for i, erow in enumerate(edf.itertuples(index=False, name=None), start=1):
                    ews.write_row(i, 0, _row(erow))
            wb.close()
        else:
            from openpyxl import Workbook
            wb = Workbook(write_only=True)
            ws = wb.create_sheet("Sheet1"); ws.append(cols)
            for _ in range(max_rows):
                try: row = next(it)
                except StopIteration: break
                ws.append(_row(row)); rows += 1
            for nm, edf in (extras or {}).items():
                ews = wb.create_sheet(str(nm)[:31])
                ews.append([str(c) for c in edf.columns])
                for erow in edf.itertuples(index=False, name=None):
                    ews.append(_row(erow))
            wb.save(target)
        written.append(target)
        if nparts > 1:
            print(f"  [save_excel] part {pi}/{nparts}: {rows:,} rows -> "
                  f"{os.path.basename(target)} ({engine})")
    if nparts > 1:
        print(f"  [save_excel] {n:,} rows > {max_rows:,} -> split into {nparts} files")
    else:
        print(f"  [save_excel] wrote {n:,} rows -> "
              f"{os.path.basename(path)} ({engine})")
    return written[0] if nparts == 1 else written

def simple_excel(data_array, file_name):
    print(f"------WRITING DATA TO EXCEL FILE {file_name} -------")
    save_excel(pd.DataFrame(data_array), file_name)
    print(f"------{file_name} EXCEL FILE SAVED -------")


def simple_excel_group_by_column(data_array, file_name, column):
    print(f"------WRITING DATA TO EXCEL FILE {file_name} -------")

    writer = pd.ExcelWriter(file_name, engine="xlsxwriter")

    df = pd.DataFrame(data_array)

    unique_groups = df[column].unique()

    for group in unique_groups:
        group_df = df[df[column] == group]
        sheet_name = f"{group}"
        group_df.to_excel(writer, sheet_name=sheet_name, index=False)

    writer._save()
    print(f"------{file_name} EXCEL FILE SAVED -------")


def simple_excel_with_key_as_sheets(data_obj, file_name):
    writer = pd.ExcelWriter(file_name, engine="xlsxwriter")
    for sheet, data_set in data_obj.items():
        df = pd.DataFrame(data_set)
        df.to_excel(writer, sheet_name=sheet, index=False)
    writer._save()


def replace_place_holder(query, placeholders, replace_values):
    data_str = json.dumps(query)
    for idx, item in enumerate(placeholders):
        if type(replace_values[idx]) is int:
            item_with_quotes = f'"{item}"'
            data_str = data_str.replace(item_with_quotes, str(replace_values[idx]))
            continue
        data_str = data_str.replace(item, replace_values[idx])
    replaced_query = json.loads(data_str)
    return replaced_query


# def load_replace_query(query_name, placeholders, replace_values):
#     with open(PATH_TO_QUERIES_JSON, "r") as queries:
#         all_queries = json.load(queries)
#     query = all_queries.get(query_name)
#     data_str = json.dumps(query)
#     for idx, item in enumerate(placeholders):
#         if type(replace_values[idx]) is int:
#             item_with_quotes = f'"{item}"'
#             data_str = data_str.replace(item_with_quotes, str(replace_values[idx]))
#             continue
#         data_str = data_str.replace(item, replace_values[idx])
#     replaced_query = json.loads(data_str)
#     return replaced_query


# def load_multi_queries(queries):
#     print(queries)
#     with open(PATH_TO_QUERIES_JSON, "r") as queries_all:
#         all_queries = json.load(queries_all)
#     return [all_queries.get(queries[i]) for i in range(len(queries))]


def today_date():
    current_datetime = datetime.datetime.now()
    time_str = current_datetime.strftime("%d_%b").upper()
    return time_str


def yesterdays_date():
    current_datetime = datetime.datetime.now()
    yesterday_datetime = current_datetime - datetime.timedelta(days=1)
    date_str = yesterday_datetime.strftime("%d_%b").upper()
    return date_str

